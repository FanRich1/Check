import io
import zipfile
from decimal import Decimal
from email.message import EmailMessage

import openpyxl
import pytest
import xlwt
from docx import Document
from PIL import Image
from reportlab.pdfgen import canvas

from shipping_billing.parsers.contracts import ParseLimits
from shipping_billing.parsers.engine import DocumentParser


def pdf_bytes(pages=1):
    stream = io.BytesIO()
    doc = canvas.Canvas(stream)
    for number in range(pages):
        doc.drawString(50, 750, f"PAGE {number + 1} Invoice TEST-001 Currency USD")
        doc.drawString(50, 700, "HANDLING CHARGE 65.00")
        doc.drawString(50, 670, "HANDLING CHARGE 35.00")
        doc.showPage()
    doc.save()
    return stream.getvalue()


def image_bytes(kind="PNG", frames=1):
    stream = io.BytesIO()
    first = Image.new("RGB", (30, 20), "white")
    kwargs = (
        {"save_all": True, "append_images": [Image.new("RGB", (30, 20), "blue") for _ in range(frames - 1)]}
        if frames > 1
        else {}
    )
    first.save(stream, format=kind, **kwargs)
    return stream.getvalue()


def docx_bytes():
    stream = io.BytesIO()
    doc = Document()
    doc.add_paragraph("Invoice DOCX-001 USD")
    table = doc.add_table(rows=3, cols=2)
    for row, values in zip(
        table.rows, [("Charge", "Amount"), ("HANDLING CHARGE", "65.00"), ("HANDLING CHARGE", "35.00")]
    ):
        for cell, value in zip(row.cells, values):
            cell.text = value
    doc.add_paragraph("After the fee table")
    doc.sections[0].header.paragraphs[0].text = "Header issuer"
    doc.save(stream)
    return stream.getvalue()


def workbook_bytes(kind="xlsx", formula=False):
    stream = io.BytesIO()
    if kind == "xlsx":
        book = openpyxl.Workbook()
        sheet = book.active
        sheet.title = "Fees"
        for row in [
            ("Charge", "Amount", "Currency"),
            ("HANDLING CHARGE", 65, "USD"),
            ("HANDLING CHARGE", 35, "USD"),
        ]:
            sheet.append(row)
        if formula:
            sheet["B4"] = "=SUM(B2:B3)"
        book.create_sheet("Hidden").sheet_state = "hidden"
        book["Hidden"]["A1"] = "Supporting fact"
        book.save(stream)
    else:
        book = xlwt.Workbook()
        sheet = book.add_sheet("Fees")
        for r, row in enumerate(
            [("Charge", "Amount", "Currency"), ("HANDLING CHARGE", 65, "USD"), ("HANDLING CHARGE", 35, "USD")]
        ):
            for c, value in enumerate(row):
                sheet.write(r, c, value)
        if formula:
            sheet.write(3, 1, xlwt.Formula("SUM(B2:B3)"))
        book.save(stream)
    return stream.getvalue()


def parse(tmp_path, data, name="sample", **kwargs):
    parser = DocumentParser(storage=tmp_path / "store", **kwargs)
    source = parser.parse_bytes(data, "test-source", name)
    return source, parser.report


def codes(source):
    return {i.code for i in source.issues}


def test_pdf_pages_and_distinct_same_name_rows(tmp_path):
    source, _ = parse(tmp_path, pdf_bytes(2), "bad-name.xlsx")
    assert source.detected_format == "pdf"
    assert source.metadata["pages"] == 2
    assert [f.locator.page for f in source.fragments] == [1, 2]
    for fragment in source.fragments:
        assert fragment.text.count("HANDLING CHARGE") == 2
        assert "65.00" in fragment.text and "35.00" in fragment.text
    assert "OCR_NOT_RUN" in codes(source)
    assert "FILE_EXTENSION_MISMATCH" in codes(source)


def test_pdf_page_provider_uses_single_page_input(tmp_path):
    from pypdf import PdfReader

    calls = []

    def provider(data, sid, page, content_type):
        assert len(PdfReader(io.BytesIO(data)).pages) == 1
        calls.append((sid, page, content_type))
        return {"mode": "test_injected", "documents": [], "text_blocks": [f"text page {page}"]}

    source, _ = parse(tmp_path, pdf_bytes(2), ocr_provider=provider)
    assert [c[1] for c in calls] == [1, 2]
    assert [f.locator.page for f in source.fragments if f.method == "specified_ocr_text"] == [1, 2]


@pytest.mark.parametrize("kind,detected", [("PNG", "png"), ("JPEG", "jpg"), ("TIFF", "tiff")])
def test_image_signatures_and_frame_inventory(tmp_path, kind, detected):
    source, _ = parse(tmp_path, image_bytes(kind), "anything.bin")
    assert source.detected_format == detected
    assert source.metadata["frames"] == 1
    assert source.metadata["page_inventory"][0]["width"] == 30
    assert "OCR_NOT_RUN" in codes(source)


def test_tiff_frames_transcoded_individually(tmp_path):
    calls = []

    def provider(data, sid, page, content_type):
        assert data.startswith(b"\x89PNG")
        assert content_type == "image/png"
        calls.append(page)
        return {
            "mode": "test_injected",
            "documents": [{"vendor_facts": {"invoiceNumber": "1"}}],
            "text_blocks": [],
        }

    source, _ = parse(tmp_path, image_bytes("TIFF", 2), ocr_provider=provider)
    assert calls == [1, 2]
    assert [o["source_page"] for o in source.ocr_results] == [1, 2]
    assert not source.fragments
    assert "OCR_NO_ORIGINAL_TEXT" in codes(source)


def test_office_rows_and_word_positions(tmp_path):
    source, _ = parse(tmp_path, docx_bytes(), "wrong.pdf")
    assert source.detected_format == "docx"
    assert source.status == "parsed"
    assert source.tables[0].rows[1:] == [["HANDLING CHARGE", "65.00"], ["HANDLING CHARGE", "35.00"]]
    texts = [f.text for f in source.fragments]
    assert texts.index("Invoice DOCX-001 USD") < texts.index("65.00") < texts.index("After the fee table")
    assert "Header issuer" in texts
    assert all(f.locator.block for f in source.fragments)


@pytest.mark.parametrize("kind", ["xls", "xlsx"])
def test_excel_actual_files_keep_same_name_rows(tmp_path, kind):
    source, _ = parse(tmp_path, workbook_bytes(kind), f"sample.{kind}")
    assert source.status == "parsed"
    assert source.detected_format == kind
    rows = [t for t in source.tables if t.locator.sheet == "Fees"]
    assert rows[1].rows == [["HANDLING CHARGE", Decimal(65), "USD"]]
    assert rows[2].rows == [["HANDLING CHARGE", Decimal(35), "USD"]]
    assert rows[2].locator.cells == "A3:C3"


@pytest.mark.parametrize("kind", ["xls", "xlsx"])
def test_formula_without_cache_is_not_zero(tmp_path, kind):
    source, _ = parse(tmp_path, workbook_bytes(kind, True), f"sample.{kind}")
    assert source.status == "pending_review"
    assert "FORMULA_CACHE_MISSING" in codes(source)
    row = next(t for t in source.tables if t.locator.sheet == "Fees" and t.locator.cells.startswith("A4"))
    assert row.rows[0][1] in (None, "")


def test_formula_cached_result_is_preserved_but_not_assumed_fresh(tmp_path):
    data = workbook_bytes(formula=True)
    buffer = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(data)) as original, zipfile.ZipFile(buffer, "w") as changed:
        for item in original.infolist():
            content = original.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                content = content.replace(b"<f>SUM(B2:B3)</f><v></v>", b"<f>SUM(B2:B3)</f><v>100</v>")
            changed.writestr(item, content)
    source, _ = parse(tmp_path, buffer.getvalue())
    assert "FORMULA_CACHE_UNVERIFIED" in codes(source)
    row = next(t for t in source.tables if t.locator.sheet == "Fees" and t.locator.cells.startswith("A4"))
    assert row.rows[0][1] == 100
    assert row.metadata["cells"][1]["formula"] == "=SUM(B2:B3)"


def test_email_alternative_octet_pdf_inline_image_and_chinese_header(tmp_path):
    email = EmailMessage()
    email["Subject"] = "账单核对"
    email["From"] = "a@example.invalid"
    email.set_content("Actual handling fee USD 50")
    email.add_alternative("<html><body><p>Actual handling fee USD 50</p></body></html>", subtype="html")
    email.add_attachment(
        pdf_bytes(), maintype="application", subtype="octet-stream", filename="../../invoice.txt"
    )
    email.add_attachment(
        image_bytes(),
        maintype="image",
        subtype="png",
        filename="signature.png",
        disposition="inline",
        cid="<logo>",
    )
    source, report = parse(tmp_path, email.as_bytes(), "sample.eml")
    assert source.metadata["headers"]["Subject"] == ["账单核对"]
    bodies = [f for f in source.fragments if f.locator.kind == "email_body"]
    assert len(bodies) == 1 and "USD 50" in bodies[0].text
    assert len(report.sources) == 3
    pdf = next(s for s in report.sources if s.detected_format == "pdf")
    image = next(s for s in report.sources if s.detected_format == "png")
    assert pdf.declared_mime == "application/octet-stream"
    assert pdf.parent_source_id == source.source_id
    assert image.metadata["inline"] is True
    assert image.metadata["classification"] == "unclassified_attachment"
    assert not (tmp_path / "invoice.txt").exists()


def test_html_only_keeps_table_cell_boundaries_and_ignores_script(tmp_path):
    email = EmailMessage()
    email["Subject"] = "HTML"
    email["From"] = "a@example.invalid"
    email.set_content(
        "<html><style>hidden</style><table><tr><td>Charge</td><td>40 USD</td></tr></table><script>evil</script></html>",
        subtype="html",
    )
    source, _ = parse(tmp_path, email.as_bytes())
    text = next(f.text for f in source.fragments if f.locator.kind == "email_body")
    assert "Charge\t40 USD" in text
    assert "evil" not in text and "hidden" not in text


def test_nested_email_parent_links_and_depth_limit(tmp_path):
    child = EmailMessage()
    child["Subject"] = "Child"
    child["From"] = "a@example.invalid"
    child.set_content("USD 20")
    root = EmailMessage()
    root["Subject"] = "Root"
    root["From"] = "a@example.invalid"
    root.set_content("See attached")
    root.add_attachment(child)
    _, report = parse(tmp_path, root.as_bytes())
    assert len(report.sources) == 2
    assert report.sources[1].parent_source_id == report.sources[0].source_id
    assert "USD 20" in "\n".join(f.text for f in report.sources[1].fragments)
    _, limited = parse(tmp_path / "limited", root.as_bytes(), limits=ParseLimits(max_depth=0))
    assert limited.sources[1].status == "error"
    assert "NESTING_LIMIT" in codes(limited.sources[1])


@pytest.mark.parametrize(
    "data", [b"%PDF-1.4\nnot a PDF", b"PK\x03\x04broken", b"\x89PNG\r\n\x1a\nbroken", b"garbage"]
)
def test_corrupt_file_has_error_entry(tmp_path, data):
    source, report = parse(tmp_path, data)
    assert source.status == "error" and len(report.sources) == 1


def test_limits(tmp_path):
    source, _ = parse(tmp_path, pdf_bytes(2), limits=ParseLimits(max_pages=1))
    assert "PAGE_LIMIT" in codes(source)
    source, _ = parse(tmp_path, workbook_bytes(), limits=ParseLimits(max_cells=3))
    assert "CELL_LIMIT" in codes(source)
    source, _ = parse(tmp_path, image_bytes(), limits=ParseLimits(max_image_pixels=10))
    assert "IMAGE_PIXEL_LIMIT" in codes(source)
    source, _ = parse(tmp_path, docx_bytes(), limits=ParseLimits(max_zip_bytes=100))
    assert "ARCHIVE_LIMIT" in codes(source)


def test_ocr_error_preserves_native_source(tmp_path):
    from shipping_billing.errors import BillingError

    def fail(*args):
        raise BillingError("OCR_HTTP_FAILED", "OCR failed.")

    source, _ = parse(tmp_path, pdf_bytes(), ocr_provider=fail)
    assert "HANDLING CHARGE 65.00" in source.fragments[0].text
    assert "OCR_HTTP_FAILED" in codes(source)
    assert source.status == "pending_review"


def test_missing_input_does_not_discard_other_sources(tmp_path):
    engine = DocumentParser(storage=tmp_path / "store")
    bad = engine.parse_file(tmp_path / "missing.pdf", "missing")
    good = engine.parse_bytes(docx_bytes(), "good", "file.docx")
    assert bad.sha256 is None and bad.status == "error"
    assert good.status == "parsed"
    assert len(engine.report.sources) == 2


def test_oversize_root_is_registered_without_reading(tmp_path):
    file = tmp_path / "file.pdf"
    file.write_bytes(pdf_bytes())
    engine = DocumentParser(storage=tmp_path / "store", limits=ParseLimits(max_file_bytes=5))
    source = engine.parse_file(file, "large")
    assert source.sha256 is None
    assert "FILE_SIZE_LIMIT" in codes(source)


def test_merged_word_cells_and_nested_tables(tmp_path):
    document = Document()
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).merge(table.cell(0, 1)).text = "Merged header"
    nested = table.cell(1, 0).add_table(rows=1, cols=2)
    nested.cell(0, 0).text = "Nested fee"
    nested.cell(0, 1).text = "12.00"
    stream = io.BytesIO()
    document.save(stream)
    source, _ = parse(tmp_path, stream.getvalue())
    assert sum(f.text == "Merged header" for f in source.fragments) == 1
    assert any(t.rows == [["Nested fee", "12.00"]] for t in source.tables)
    assert any(t.metadata["merged_cells"] for t in source.tables)


def test_doc_signature_and_conversion_dispatch(tmp_path):
    from pathlib import Path

    data = (Path(__file__).parents[1] / "fixtures/parsing/fees.doc").read_bytes()
    calls = []

    def converter(blob, storage):
        calls.append(blob)
        return docx_bytes()

    source, _ = parse(tmp_path, data, "fees.doc", doc_converter=converter)
    assert source.detected_format == "doc" and source.status == "parsed"
    assert calls == [data]
    assert source.metadata["conversion"]["output_sha256"]
    assert source.tables[0].rows[1][1] == "65.00"


def test_local_ocr_keeps_regions_and_is_not_remote_success(tmp_path):
    class Local:
        metadata = {"engine": "test_injected"}

        def __call__(self, image):
            return [{"polygon": [[0, 0], [20, 0], [20, 10], [0, 10]], "text": "USD 40", "confidence": 0.7}]

    source, _ = parse(tmp_path, image_bytes(), local_ocr=Local())
    assert source.fragments[0].method == "local_ocr_supplement"
    assert "polygon_pixels" in source.fragments[0].locator.region
    assert "OCR_NOT_RUN" in codes(source) and "LOCAL_OCR_LOW_CONFIDENCE" in codes(source)
    assert source.ocr_results == []


def test_scanned_pdf_keeps_blank_native_page_and_local_evidence(tmp_path):
    from pathlib import Path

    class Local:
        metadata = {"engine": "test_injected"}

        def __call__(self, image):
            return [{"polygon": [[0, 0], [20, 0], [20, 10], [0, 10]], "text": "USD 100", "confidence": 0.99}]

    data = (Path(__file__).parents[1] / "fixtures/parsing/fees-scan.pdf").read_bytes()
    source, _ = parse(tmp_path, data, local_ocr=Local())
    assert source.metadata["pages"] == 2
    assert len(source.metadata["local_ocr_pages"]) == 2
    assert all(not f.text for f in source.fragments if f.method == "pdf_native_text")
    assert [f.locator.page for f in source.fragments if f.method == "local_ocr_supplement"] == [1, 2]


def test_vendor_merged_rows_are_flagged_without_overwriting(tmp_path):
    def provider(*args):
        return {
            "mode": "test_injected",
            "text_blocks": [],
            "documents": [
                {
                    "vendor_facts": {
                        "invoiceNumber": "MISMATCH-001",
                        "chargeList": [{"chargeItem": "HANDLING CHARGE", "amount": "100.00"}],
                    }
                }
            ],
        }

    source, _ = parse(tmp_path, pdf_bytes(), ocr_provider=provider)
    assert "OCR_POSSIBLE_MERGED_FEE_ROWS" in codes(source)
    assert "OCR_IDENTIFIER_NOT_CORROBORATED" in codes(source)
    assert source.ocr_results[0]["documents"][0]["vendor_facts"]["chargeList"][0]["amount"] == "100.00"
    assert source.fragments[0].text.count("HANDLING CHARGE") == 2


def test_base64_nested_email_is_decoded(tmp_path):
    import base64

    inner = b"From: a@example.invalid\r\nSubject: Nested\r\n\r\nfee USD 20\r\n"
    data = (
        b"From: a@example.invalid\r\nSubject: outer\r\nMIME-Version: 1.0\r\nContent-Type: message/rfc822\r\nContent-Transfer-Encoding: base64\r\n\r\n"
        + base64.b64encode(inner)
    )
    _, report = parse(tmp_path, data)
    assert len(report.sources) == 2
    assert "fee USD 20" in "\n".join(f.text for f in report.sources[1].fragments)


def test_encrypted_pdf_is_explicit_error(tmp_path):
    from pypdf import PdfReader, PdfWriter

    writer = PdfWriter()
    writer.append(PdfReader(io.BytesIO(pdf_bytes())))
    writer.encrypt("secret")
    stream = io.BytesIO()
    writer.write(stream)
    source, _ = parse(tmp_path, stream.getvalue())
    assert "ENCRYPTED_FILE" in codes(source) and source.status == "error"


def test_email_source_limit_is_visible(tmp_path):
    root = EmailMessage()
    root["From"] = "a@example.invalid"
    root["Subject"] = "many"
    root.set_content("files")
    for i in range(3):
        root.add_attachment(pdf_bytes(), maintype="application", subtype="pdf", filename=f"{i}.pdf")
    source, report = parse(tmp_path, root.as_bytes(), limits=ParseLimits(max_sources=2))
    assert len(report.sources) == 2
    assert source.metadata["unexpanded_attachment_count"] == 2
    assert source.status == "error"


def test_fee_name_prefix_is_not_a_duplicate_row(tmp_path):
    stream = io.BytesIO()
    doc = canvas.Canvas(stream)
    doc.drawString(40, 700, "WHARFAGE ADMINISTRATION FEE 5.00")
    doc.drawString(40, 670, "WHARFAGE 21.90")
    doc.save()

    def provider(*args):
        return {
            "mode": "test_injected",
            "text_blocks": [],
            "documents": [
                {
                    "vendor_facts": {
                        "chargeList": [
                            {"chargeItem": "WHARFAGE ADMINISTRATION FEE", "amount": "5.00"},
                            {"chargeItem": "WHARFAGE", "amount": "21.90"},
                        ]
                    }
                }
            ],
        }

    source, _ = parse(tmp_path, stream.getvalue(), ocr_provider=provider)
    assert "OCR_POSSIBLE_MERGED_FEE_ROWS" not in codes(source)
