"""Office extraction preserves rows and positions; never evaluates workbook formulas."""

import io
import struct
from datetime import date, datetime, time
from decimal import Decimal

import olefile
import openpyxl
import xlrd
from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from openpyxl.utils import get_column_letter

from shipping_billing.errors import BillingError
from shipping_billing.models import Locator
from shipping_billing.parsers.contracts import ParsedTable


def scalar(value):
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if isinstance(value, float):
        return Decimal(str(value))
    if value is None or isinstance(value, (str, int, bool, Decimal)):
        return value
    return str(value)


def parse_xlsx(data, source, limits):
    book = openpyxl.load_workbook(io.BytesIO(data), data_only=False, keep_links=False)
    cached = openpyxl.load_workbook(io.BytesIO(data), data_only=True, keep_links=False)
    try:
        total = sum(s.max_row * s.max_column for s in book.worksheets)
        if total > limits.max_cells:
            raise BillingError("CELL_LIMIT", "Workbook exceeds configured cell limit.")
        source.metadata["sheets"] = []
        for sheet in book.worksheets:
            source.metadata["sheets"].append(
                {
                    "name": sheet.title,
                    "state": sheet.sheet_state,
                    "merged_ranges": [str(r) for r in sheet.merged_cells.ranges],
                }
            )
            for row in sheet.iter_rows():
                if all(c.value is None for c in row):
                    continue
                values, details = [], []
                for cell in row:
                    value = cell.value
                    detail = {
                        "coordinate": cell.coordinate,
                        "data_type": cell.data_type,
                        "number_format": cell.number_format,
                        "formula": None,
                        "hidden_column": bool(sheet.column_dimensions[cell.column_letter].hidden)
                        if hasattr(cell, "column_letter")
                        else False,
                    }
                    if cell.data_type == "f":
                        detail["formula"] = (
                            value if isinstance(value, str) else str(getattr(value, "text", value))
                        )
                        value = cached[sheet.title][cell.coordinate].value
                        detail["cached_value"] = scalar(value)
                        detail["cache_verified_fresh"] = False
                        if value is None:
                            source.issue(
                                "FORMULA_CACHE_MISSING",
                                "A formula has no cached result; no recalculation performed.",
                            )
                        else:
                            source.issue(
                                "FORMULA_CACHE_UNVERIFIED",
                                "Formula result is a stored cache; freshness requires verification.",
                            )
                    if cell.data_type == "e" or cached[sheet.title][cell.coordinate].data_type == "e":
                        source.issue("SPREADSHEET_CELL_ERROR", "Workbook contains an Excel error value.")
                    values.append(scalar(value))
                    details.append(detail)
                rownum = row[0].row
                address = f"A{rownum}:{get_column_letter(len(row))}{rownum}"
                locator = Locator(kind="sheet_cells", sheet=sheet.title, cells=address)
                source.tables.append(
                    ParsedTable(
                        table_id=f"{source.source_id}:{sheet.title}:{rownum}",
                        locator=locator,
                        rows=[values],
                        metadata={"cells": details, "hidden_row": bool(sheet.row_dimensions[rownum].hidden)},
                    )
                )
                source.fragment("\t".join("" if x is None else str(x) for x in values), locator, "xlsx_cells")
        calculation = book.calculation
        source.metadata["calculation"] = {
            "mode": getattr(calculation, "calcMode", None),
            "full_calc_on_load": getattr(calculation, "fullCalcOnLoad", None),
        }
    finally:
        book.close()
        cached.close()


def xls_formula_cells(data):
    """Read BIFF8 formula locations only; expression tokens are not evaluated."""
    with olefile.OleFileIO(io.BytesIO(data)) as archive:
        raw = archive.openstream("Workbook" if archive.exists("Workbook") else "Book").read()
    offsets, records = [], []
    pos = 0
    while pos + 4 <= len(raw):
        kind, size = struct.unpack_from("<HH", raw, pos)
        payload = raw[pos + 4 : pos + 4 + size]
        if len(payload) != size:
            break
        if kind == 0x85 and len(payload) >= 8:
            offset = struct.unpack_from("<I", payload)[0]
            length, flags = payload[6:8]
            name = payload[8 : 8 + length * (2 if flags & 1 else 1)].decode(
                "utf-16-le" if flags & 1 else "latin1"
            )
            offsets.append((offset, name))
        if kind == 0x06 and len(payload) >= 14:
            records.append((pos, struct.unpack_from("<HH", payload)))
        pos += size + 4
    result = set()
    for position, (row, col) in records:
        earlier = [(offset, name) for offset, name in offsets if offset <= position]
        if earlier:
            result.add((max(earlier)[1], row, col))
    return result


def parse_xls(data, source, limits):
    formulas = xls_formula_cells(data)
    book = xlrd.open_workbook(file_contents=data, formatting_info=True, on_demand=True)
    try:
        if (
            sum(book.sheet_by_index(i).nrows * book.sheet_by_index(i).ncols for i in range(book.nsheets))
            > limits.max_cells
        ):
            raise BillingError("CELL_LIMIT", "Workbook exceeds configured cell limit.")
        source.metadata["sheets"] = []
        source.metadata["date_mode"] = book.datemode
        for sheet in book.sheets():
            source.metadata["sheets"].append(
                {"name": sheet.name, "visibility": sheet.visibility, "merged_ranges": sheet.merged_cells}
            )
            for rownum in range(sheet.nrows):
                cells = sheet.row(rownum)
                if all(c.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK) for c in cells):
                    continue
                values, details = [], []
                for col, cell in enumerate(cells):
                    value = cell.value
                    if cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                        value = None
                    elif cell.ctype == xlrd.XL_CELL_DATE:
                        value = xlrd.xldate_as_datetime(value, book.datemode).isoformat()
                    elif cell.ctype == xlrd.XL_CELL_BOOLEAN:
                        value = bool(value)
                    elif cell.ctype == xlrd.XL_CELL_ERROR:
                        value = xlrd.error_text_from_code.get(value, "UNKNOWN_ERROR")
                        source.issue("SPREADSHEET_CELL_ERROR", "Workbook contains an Excel error value.")
                    is_formula = (sheet.name, rownum, col) in formulas
                    if is_formula:
                        source.issue(
                            "FORMULA_CACHE_MISSING" if value in (None, "") else "FORMULA_CACHE_UNVERIFIED",
                            "XLS formula tokens were detected; cached values are not recalculated or freshness-verified.",
                        )
                    details.append(
                        {
                            "coordinate": f"{get_column_letter(col + 1)}{rownum + 1}",
                            "data_type": cell.ctype,
                            "formula_present": is_formula,
                            "formula_text_available": False,
                            "cache_verified_fresh": False if is_formula else None,
                        }
                    )
                    values.append(scalar(value))
                locator = Locator(
                    kind="sheet_cells",
                    sheet=sheet.name,
                    cells=f"A{rownum + 1}:{get_column_letter(sheet.ncols)}{rownum + 1}",
                )
                source.tables.append(
                    ParsedTable(
                        table_id=f"{source.source_id}:{sheet.name}:{rownum + 1}",
                        locator=locator,
                        rows=[values],
                        metadata={"cells": details},
                    )
                )
                source.fragment("\t".join("" if x is None else str(x) for x in values), locator, "xls_cells")
    finally:
        book.release_resources()


def parse_docx(data, source, limits, visit_child):
    doc = Document(io.BytesIO(data))
    cell_count = 0

    def blocks(container, prefix):
        nonlocal cell_count
        for i, block in enumerate(container.iter_inner_content(), 1):
            name = f"{prefix}/block-{i}"
            if isinstance(block, Paragraph):
                source.fragment(block.text, Locator(kind="word_block", block=name), "docx_paragraph")
            elif isinstance(block, Table):
                rows, merged, seen = [], [], {}
                for r, row in enumerate(block.rows, 1):
                    values = []
                    for c, cell in enumerate(row.cells, 1):
                        cell_count += 1
                        if cell_count > limits.max_cells:
                            raise BillingError("CELL_LIMIT", "Word tables exceed configured cell limit.")
                        location = f"{name}/R{r}C{c}"
                        if cell._tc in seen:
                            values.append(None)
                            merged.append({"cell": location, "same_as": seen[cell._tc]})
                        else:
                            seen[cell._tc] = location
                            values.append(cell.text)
                            blocks(cell, location)
                    rows.append(values)
                source.tables.append(
                    ParsedTable(
                        table_id=f"{source.source_id}:{name}",
                        locator=Locator(kind="word_block", block=name),
                        rows=rows,
                        metadata={"merged_cells": merged},
                    )
                )

    blocks(doc, "body")
    seen_parts = set()
    for index, section in enumerate(doc.sections, 1):
        for name in (
            "header",
            "footer",
            "first_page_header",
            "first_page_footer",
            "even_page_header",
            "even_page_footer",
        ):
            part = getattr(section, name)
            if str(part.part.partname) not in seen_parts:
                seen_parts.add(str(part.part.partname))
                blocks(part, f"section-{index}/{name}")
    # Keep drawings as child sources instead of silently dropping scanned content.
    for i, part in enumerate(doc.part.package.parts, 1):
        if part.content_type.startswith("image/") and str(part.partname).startswith("/word/media/"):
            visit_child(
                part.blob,
                str(part.partname).rsplit("/", 1)[-1],
                part.content_type,
                f"image-{i}",
                {"word_part": str(part.partname), "classification": "embedded_image"},
            )
    tags = {element.tag.rsplit("}", 1)[-1] for element in doc.element.iter()}
    if tags & {"ins", "del", "txbxContent", "altChunk", "object"}:
        source.issue(
            "WORD_COMPLEX_CONTENT", "Tracked edits, text boxes or embedded objects require additional review."
        )
