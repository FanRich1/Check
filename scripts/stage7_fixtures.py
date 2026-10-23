"""Synthetic documents and a clearly injected model for deterministic workflow tests.

The injected adapter reads source facts; no operation labels or expected records are supplied.
"""

import re
from decimal import Decimal
from email.message import EmailMessage

import simplejson

from shipping_billing.errors import BillingError
from shipping_billing.extraction.contracts import ExtractionDraft
from shipping_billing.pipeline.intents import IntentDraft

HEADER = ["Type: Invoice", "Invoice: INV-E2E-001", "Issuer: Supplier Ltd", "Bill to: Buyer Ltd"]


def content(amount=40, *, note=None, rows=None, currency="USD"):
    lines = HEADER + [f"Handling Fee | {currency} | {amount} | PP"] if rows is None else HEADER + rows
    if note:
        lines.append(note)
    return lines


def eml(path, amount=40, *, note=None, rows=None, currency="USD", attachment=None):
    message = EmailMessage()
    message["From"] = "supplier@example.invalid"
    message["To"] = "buyer@example.invalid"
    message["Subject"] = "Synthetic workflow invoice"
    message["Date"] = "Fri, 23 Oct 2026 08:00:00 +0000"
    message["Message-ID"] = "<stage7-synthetic@example.invalid>"
    message.set_content("\n".join(content(amount, note=note, rows=rows, currency=currency)))
    if attachment:
        message.add_attachment(
            attachment.read_bytes(), maintype="application", subtype="octet-stream", filename=attachment.name
        )
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(message.as_bytes())
    return path


def documents(folder):
    import xlwt
    from docx import Document
    from openpyxl import Workbook
    from PIL import Image, ImageDraw
    from reportlab.pdfgen.canvas import Canvas

    folder.mkdir(parents=True, exist_ok=True)
    lines = content()
    result = {"eml": eml(folder / "invoice.eml")}
    book = Workbook()
    for text in lines:
        book.active.append([text])
    book.save(folder / "invoice.xlsx")
    result["xlsx"] = folder / "invoice.xlsx"
    old = xlwt.Workbook()
    sheet = old.add_sheet("Invoice")
    for row, text in enumerate(lines):
        sheet.write(row, 0, text)
    old.save(str(folder / "invoice.xls"))
    result["xls"] = folder / "invoice.xls"
    doc = Document()
    for text in lines:
        doc.add_paragraph(text)
    doc.save(folder / "invoice.docx")
    result["docx"] = folder / "invoice.docx"
    pdf = Canvas(str(folder / "invoice.pdf"), invariant=1)
    for i, text in enumerate(lines):
        pdf.drawString(40, 800 - 22 * i, text)
    pdf.save()
    result["pdf"] = folder / "invoice.pdf"
    image = Image.new("RGB", (720, 180), "white")
    draw = ImageDraw.Draw(image)
    for i, text in enumerate(lines):
        draw.text((15, 10 + 25 * i), text, fill="black")
    for fmt in ("png", "jpg", "tiff"):
        path = folder / ("invoice." + fmt)
        image.save(path)
        result[fmt] = path
    return result


def fake_ocr(content_bytes, sid, page, content_type):
    return {"mode": "test_injected", "network_calls": 0, "documents": [], "text_blocks": content()}


class AuthoredModel:
    """A test double, never a live provider or a claimed extraction accuracy evaluation."""

    def __init__(self, *, error=None, mutate_intent=None):
        self.error, self.mutate_intent = error, mutate_intent
        self.calls = []

    def extract(self, text, result_type, **kwargs):
        payload = simplejson.loads(text, use_decimal=True)
        self.calls.append(result_type.__name__)
        if self.error:
            raise BillingError(self.error, "Injected model failure.")
        if result_type is IntentDraft:
            return self.intents(payload), {"mode": "test_injected"}
        assert result_type is ExtractionDraft
        units = payload["target_units"]

        def find(pattern):
            return next(((u, m) for u in units if (m := re.search(pattern, u["text"]))), None)

        def header(pattern):
            hit = find(pattern)
            return (
                (hit[1].group(1), [dict(unit_id=hit[0]["unit_id"], quote=hit[0]["text"])])
                if hit
                else (None, [])
            )

        values, evidence = {}, {}
        for field, pattern in [
            ("document_type", r"Type: (Invoice)"),
            ("reference_no", r"Invoice: (INV-[\w-]+)"),
            ("issuer", r"Issuer: (Supplier Ltd)"),
            ("recipient", r"Bill to: (Buyer Ltd)"),
        ]:
            values[field], evidence[field] = header(pattern)
        fees = []
        has_native = any(u["method"] == "pdf_native_text" and " | " in u["text"] for u in units)
        for unit in units:
            if has_native and unit["method"] == "specified_ocr_text":
                continue
            match = re.search(
                r"(Handling Fee|Storage Fee)\s*\|\s*(USD|EUR|\$)\s*\|\s*([\d.]+)\s*\|\s*(PP|CC)", unit["text"]
            )
            if not match:
                continue
            ref = dict(unit_id=unit["unit_id"], quote=unit["text"])
            fees.append(
                dict(
                    charge_item_raw=match[1],
                    currency_raw=match[2],
                    amount=Decimal(match[3]),
                    quantity=None,
                    unit_price=None,
                    payment_term_raw=match[4],
                    evidence={
                        k: [ref] for k in ("charge_item_raw", "currency_raw", "amount", "payment_term_raw")
                    },
                )
            )
        documents = []
        if fees:
            documents = [
                dict(**values, document_date_raw=None, shipping_info={}, evidence=evidence, fees=fees)
            ]
        result = ExtractionDraft.model_validate(dict(documents=documents, observations=[]))
        if kwargs.get("response_sink"):
            kwargs["response_sink"](simplejson.dumps(result.model_dump(mode="python"), use_decimal=True))
        return result, {"mode": "test_injected"}

    def intents(self, payload):
        directives = []
        for unit in payload["units"]:
            match = re.search(r"Update amount from (\d+) to (\d+)", unit["text"])
            if not match:
                continue
            target = payload["lines"][0]["source_line_id"]
            directive = dict(
                source_line_id=target,
                record_id=None,
                selector_citations=[],
                edits=[
                    dict(
                        path="/charge_details/amount",
                        action="set",
                        expected_old_stated=True,
                        expected_old=int(match[1]),
                        value=int(match[2]),
                        citations=[dict(unit_id=unit["unit_id"], quote=unit["text"])],
                    )
                ],
            )
            directives.append(directive)
        data = {"directives": directives}
        if self.mutate_intent:
            self.mutate_intent(data, payload)
        return IntentDraft.model_validate(data)
